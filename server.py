from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import cgi,json,re,shutil,uuid
ROOT=Path(__file__).resolve().parent; UPLOADS=ROOT/'uploads'; UPLOADS.mkdir(exist_ok=True); ALLOWED={'.pdf','.docx','.txt','.md'}
def text_of(p):
 s=p.suffix.lower()
 if s in {'.txt','.md'}: return p.read_text(encoding='utf-8',errors='ignore')
 if s=='.pdf':
  from pypdf import PdfReader
  return '\n'.join(x.extract_text() or '' for x in PdfReader(str(p)).pages)
 if s=='.docx':
  from docx import Document
  d=Document(str(p)); return '\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(' '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
 raise ValueError('仅支持 PDF、Word、TXT、MD 文件')
def find(text,ps,default='未识别'):
 for p in ps:
  m=re.search(p,text,re.I)
  if m:return m.group(1).strip()
 return default
def parse(text,name):
 compact=re.sub(r'\s+',' ',text); keys=['sql','用户研究','数据分析','产品','Excel','Axure','Python','项目管理']; skills=[x.upper() if x=='sql' else x for x in keys if re.search(x,text,re.I)]; score=min(98,45+sum(v for k,v in {'sql':22,'用户研究':18,'产品':18,'数据分析':16,'本科':12,'实习':14}.items() if re.search(k,text,re.I)))
 return {'id':uuid.uuid4().hex[:8],'name':find(compact,[r'姓名[:： ]+([\u4e00-\u9fa5]{2,4})',r'^([\u4e00-\u9fa5]{2,4})'],Path(name).stem),'school':find(compact,[r'([\u4e00-\u9fa5]{2,20}(?:大学|学院))']),'education':'本科及以上' if re.search(r'本科|硕士|研究生',text) else '未识别','skills':skills,'score':score,'status':'强匹配' if score>=78 else '待复核','filename':name,'text_preview':compact[:500],'risks':['建议验证 SQL 实战深度' if 'sql' in compact.lower() else '未识别到 SQL','到岗时间需面试确认'],'questions':['请用 STAR 法则说明你如何通过数据推动一次产品改进？','请具体介绍一段用户研究或需求分析经历。']}
def reply(h,obj,status=200):
 b=json.dumps(obj,ensure_ascii=False).encode(); h.send_response(status); h.send_header('Content-Type','application/json; charset=utf-8'); h.send_header('Content-Length',str(len(b))); h.end_headers(); h.wfile.write(b)
class Handler(BaseHTTPRequestHandler):
 def log_message(self,*a):pass
 def do_GET(self):
  path=urlparse(self.path).path
  if path=='/api/health':return reply(self,{'ok':True,'service':'AI-HR-System'})
  p=ROOT/('index.html' if path=='/' else path.lstrip('/'))
  if p.is_file() and ROOT in p.resolve().parents:
   b=p.read_bytes(); self.send_response(200); self.send_header('Content-Type','text/html; charset=utf-8' if p.suffix=='.html' else 'text/plain; charset=utf-8'); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b); return
  self.send_error(404)
 def do_POST(self):
  if urlparse(self.path).path!='/api/resumes/upload':return reply(self,{'error':'Not found'},404)
  try:
   f=cgi.FieldStorage(fp=self.rfile,headers=self.headers,environ={'REQUEST_METHOD':'POST','CONTENT_TYPE':self.headers.get('Content-Type','')}); items=f['files'] if 'files' in f else []; items=items if isinstance(items,list) else [items]; out=[]
   for item in items:
    name=Path(item.filename or '').name
    if Path(name).suffix.lower() not in ALLOWED:raise ValueError(f'不支持的文件类型: {name}')
    p=UPLOADS/f'{uuid.uuid4().hex}_{name}'; item.file.seek(0)
    with p.open('wb') as w:shutil.copyfileobj(item.file,w)
    try:out.append(parse(text_of(p),name))
    finally:p.unlink(missing_ok=True)
   return reply(self,{'count':len(out),'candidates':out})
  except Exception as e:return reply(self,{'error':str(e)},400)
if __name__=='__main__':ThreadingHTTPServer(('127.0.0.1',8787),Handler).serve_forever()
